"""v0.6结果表与原始数值轨迹的独立全量核验程序。
不导入结果导出器或其采样函数。"""
from pathlib import Path
from bisect import bisect_right
from dataclasses import asdict
from datetime import datetime,timezone
import gzip,hashlib,json,math
import openpyxl
import numpy as np
from model import Inputs
from run_all import configurations

ROOT=Path(__file__).resolve().parent
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

# 独立按当前物理半径插值轨迹，域外结果保持为空且不调用导出采样器
def interp(snapshot,field,distance):
    radius=snapshot['radius_cm']
    if distance>radius+1e-10:return None
    if 'sample_distances_cm' in snapshot:
        grid=snapshot['sample_distances_cm']
    else:grid=[x*radius for x in snapshot['x']]
    values=snapshot[field]
    if distance>=grid[-1]:return values[-1]
    j=max(0,bisect_right(grid,distance)-1)
    return values[j]+(values[j+1]-values[j])*(distance-grid[j])/(grid[j+1]-grid[j])

# ==================== 结果表与原始轨迹独立核验 ====================
def main():
    report={'status':'running','sheets':[],'numeric_values_compared':0,
            'max_state_rounding_error':0.,'source_sha256':sha(ROOT/'model.py')}
    inp=Inputs()
    # 核对四份Excel对应的求解轨迹、模型配置及输入来源
    for case,numbers in [('q1_final',[1]),('q2_full',[2,3]),('q4_final',[4])]:
        with gzip.open(ROOT/'runs'/f'{case}.json.gz','rt',encoding='utf-8') as f:run=json.load(f)
        assert run['source_hash']==sha(ROOT/'model.py')
        assert run['config']==asdict(configurations()[case])
        for name,digest in run['input_sha256'].items():assert sha(ROOT.parent/'附件'/name)==digest
        # 核验严格阈值跨越、事件区间宽度及终点右端时刻
        if run['endpoint']:
            e=run['endpoint']
            assert 0<e['right_s']-e['left_s']<=.020001
            assert e['left_max_C']>=.15>e['right_max_C']
            assert run['final']['t']==e['right_s']
        for number in numbers:
            states=run['short'] if number<=2 else run['long']
            step=1 if number<=2 else 60
            end=run['final']['t']
            # 按题给秒/分钟间隔构造输出合同，并追加实际非整点终点
            times=list(range(0,int(end)+1,step))
            if times[-1]!=end:times.append(end)
            assert [s['t'] for s in states]==times
            book=ROOT/f'result{number}.xlsx'
            wb=openpyxl.load_workbook(book,read_only=True,data_only=False)
            names=['温度','水分浓度'] if number<=2 else ['Sheet1']
            assert wb.sheetnames==names
            for title,field in zip(names,['T','C'] if number<=2 else ['C']):
                ws=wb[title]
                rows=iter(ws.iter_rows())
                header=next(rows)
                assert [c.value for c in header[1:22]]==[i/10 for i in range(21)]
                count=1
                # 逐时间层核对物理距离采样、表面位置和域外空值
                for snap in states:
                    row=next(rows)
                    assert len(row)==(24 if number==4 else 22)
                    assert abs(row[0].value-snap['t'])<1e-8
                    if number==4:
                        assert abs(snap['radius_cm']-100*inp.radius(snap['t'],True))<1e-12
                    for j in range(1,len(row)):
                        expected=interp(snap,field,(j-1)/10) if j<=21 else (snap[field][-1] if j==22 else snap['radius_cm'])
                        cell=row[j];value=cell.value
                        # 分别检查域外留空与域内有限数值、四位小数及舍入误差
                        if expected is None:
                            assert value is None,(number,count,j,value)
                        else:
                            assert cell.data_type!='f'
                            assert isinstance(value,(int,float)) and math.isfinite(value)
                            assert cell.number_format=='0.0000'
                            delta=abs(value-expected)
                            assert delta<=.00005000001,(number,count,j,value,expected)
                            assert abs(value*1e4-round(value*1e4))<1e-6
                            report['max_state_rounding_error']=max(report['max_state_rounding_error'],delta)
                            report['numeric_values_compared']+=1
                    count+=1
                assert next(rows,None) is None
                report['sheets'].append({'file':book.name,'sheet':title,'rows':count,
                    'columns':24 if number==4 else 22,'last_time_s':end})
                print(number,title,count,'passed',flush=True)
            wb.close()
            assert sha(book)==sha(ROOT/'results'/book.name)
        for snap in run['long']:
            for field in ('C','T'):assert np.isfinite(snap[field]).all()
            assert min(snap['C'])>0 and max(snap['C'])<=2.55+1e-10
            assert -1e-12<=snap['wet_fraction']<=1+1e-12
        del run
    # ==================== 全量验证记录与文件指纹输出 ====================
    report.update(status='passed',completed_utc=datetime.now(timezone.utc).isoformat(),
        verifier_sha256=sha(Path(__file__)),
        xlsx_sha256={f'result{i}.xlsx':sha(ROOT/f'result{i}.xlsx') for i in range(1,5)})
    (ROOT/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False),flush=True)

if __name__=='__main__':main()
